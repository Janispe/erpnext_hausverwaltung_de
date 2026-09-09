"""Deterministischer XLSX-Inhalt aus dem gespeicherten Meldungsstand.

Text wird explizit als Text geschrieben, insbesondere auch Texte mit =/+/@.
Der Export führt weder Formeln aus Vorlagen aus noch liest er aktuelle ERP-Werte nach.
"""

from datetime import date, datetime
from io import BytesIO
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_schema import json_object, typed_value

EXPORT_VERSION = 1


def _extras(definitions, scope):
	return [d for d in definitions if d["bereich"] == scope]


def _label(d):
	return "Zusatz: " + d["bezeichnung"] + (f" ({d['einheit']})" if d.get("einheit") else "")


def _value(d, values):
	v = typed_value(d, values.get(d["schluessel"]))
	if d["feldtyp"] == "Date" and v:
		return date.fromisoformat(v)
	return v


def _date(value):
	if not value:
		return None
	return value if isinstance(value, (datetime, date)) else date.fromisoformat(str(value)[:10])


def build_xlsx(doc):
	wb = Workbook()
	wb.remove(wb.active)
	definitions = doc.definitions()
	snapshot = json_object(doc.vorlage_snapshot)
	state = "ENTWURF: noch nicht freigegeben" if int(doc.docstatus) == 0 else "Freigegebener Meldungsstand"

	def sheet(name, headers, rows, widths=None):
		ws = wb.create_sheet(name)
		ws.append([f"Heizkostenmeldung {doc.name}: {state}"])
		ws.append([str(doc.immobilie), _date(doc.von), _date(doc.bis), str(doc.liegenschaftsnummer or "")])
		ws.append([])
		ws.append(headers)
		for row in rows:
			ws.append(row)
		for row in ws:
			for cell in row:
				if isinstance(cell.value, str):
					cell.data_type = "s"
				cell.font = Font(name="Arial", size=11)
				cell.alignment = Alignment(vertical="top", wrap_text=True)
				if isinstance(cell.value, (datetime, date)):
					cell.number_format = "DD.MM.YYYY"
				elif isinstance(cell.value, (int, float)):
					cell.number_format = "#,##0.00"
		for cell in ws[4]:
			cell.fill = PatternFill("solid", fgColor="233C51")
			cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
		ws["A1"].font = Font(name="Arial", size=14, bold=True)
		ws.row_dimensions[1].height = 40
		ws.row_dimensions[4].height = 42
		for index in range(1, ws.max_column + 1):
			from openpyxl.utils import get_column_letter

			ws.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index, 24)
		ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, ws.max_column))
		for row in ws.iter_rows(min_row=5):
			lines = max(
				sum(
					max(1, ceil(len(part) / max(8, ws.column_dimensions[cell.column_letter].width - 2)))
					for part in str(cell.value or "").splitlines()
				)
				for cell in row
			)
			ws.row_dimensions[row[0].row].height = max(22, min(400, lines * 15 + 6))
		ws.freeze_panes = "C5" if name == "Nutzer" else "A5"
		ws.sheet_view.showGridLines = False
		ws.print_title_rows = "1:4"
		ws.sheet_properties.pageSetUpPr.fitToPage = True
		ws.page_setup.orientation = "landscape"
		ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws.max_column > 8 else ws.PAPERSIZE_A4
		ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
		ws.print_options.horizontalCentered = True
		return ws

	meta = [
		["Meldung", doc.name],
		["Immobilie", doc.immobilie],
		["Objektadresse", doc.objektadresse],
		["Liegenschaftsnummer", doc.liegenschaftsnummer],
		["Wärmedienst", doc.waermedienst],
		["Vorlagenversion", snapshot.get("name")],
		["Vorlagenbezeichnung", snapshot.get("bezeichnung")],
		["ERP-Datenstand", str(doc.datenstand or "Noch nicht geladen")],
		["Exportformat-Version", EXPORT_VERSION],
		["Status", state],
		[
			"Vorauszahlungen",
			"Meldebetrag je Mietvertrag; ERP-IST/SOLL separat zur Prüfung. Keine Personenzusammenfassung.",
		],
		["Notizen", doc.notizen],
	]
	sheet("Meldung", ["Angabe", "Wert"], meta, {1: 35, 2: 85})
	fields = _extras(definitions, "Nutzer")
	nutzer = []
	for r in doc.nutzer:
		values = json_object(r.zusatzwerte_json)
		nutzer.append(
			[
				r.nutzernummer,
				r.wohnung,
				r.typ,
				r.mietername,
				_date(r.von),
				_date(r.bis),
				r.heizflaeche if r.flaeche_bestaetigt else None,
				r.vorauszahlung_meldung if r.vorauszahlung_bestaetigt or r.typ == "Leerstand" else None,
				r.vorauszahlung_ist,
				r.vorauszahlung_soll,
				r.wohnflaeche,
				*[_value(d, values) for d in fields],
				r.pruefhinweise,
				r.pruefnotiz,
				r.mietvertrag,
			]
		)
	ws = sheet(
		"Nutzer",
		[
			"Nutzernummer",
			"Wohnung",
			"Nutzung",
			"Mieter",
			"Von",
			"Bis",
			"Heizfläche (m²)",
			"HK-Vorauszahlung Meldung (€)",
			"HK-IST ERP (€)",
			"HK-SOLL ERP (€)",
			"Wohnfläche ERP (m²)",
			*[_label(d) for d in fields],
			"Prüfhinweise",
			"Klärung",
			"Mietvertrag",
		],
		nutzer,
		{2: 32, 4: 35},
	)
	ws.auto_filter.ref = f"A4:{ws.cell(ws.max_row, ws.max_column).coordinate}"
	for row in range(5, ws.max_row + 1):
		ws.cell(row, 7).number_format = ws.cell(row, 11).number_format = "#,##0.000"
	fields = _extras(definitions, "Brennstoff")
	quantities = sum(float(r.menge or 0) for r in doc.lieferungen)
	rows = [
		[
			"Anfangsbestand",
			_date(doc.von),
			doc.anfangsbestand if doc.bestaende_bestaetigt else None,
			doc.anfangswert if doc.bestaende_bestaetigt else None,
		],
	]
	for r in doc.lieferungen:
		rows.append(
			[
				"Zukauf",
				_date(r.datum),
				r.menge,
				r.betrag if r.betrag_bestaetigt else None,
				r.eingangsrechnung,
				r.bemerkung,
				*[_value(d, json_object(r.zusatzwerte_json)) for d in fields],
			]
		)
	rows.extend(
		[
			["Abzüge / Gutschriften", None, None, doc.abzuege if doc.brennstoff_vollstaendig else None],
			[
				"Endbestand",
				_date(doc.bis),
				doc.endbestand if doc.bestaende_bestaetigt else None,
				None if doc.endwert_durch_messdienst or not doc.bestaende_bestaetigt else doc.endwert,
				None,
				"Wert ermittelt der Messdienst" if doc.endwert_durch_messdienst else "",
			],
			[
				"Verbrauch",
				None,
				float(doc.anfangsbestand or 0) + quantities - float(doc.endbestand or 0)
				if doc.bestaende_bestaetigt and doc.brennstoff_vollstaendig
				else None,
			],
		]
	)
	sheet(
		"Heizoel",
		[
			"Vorgang",
			"Datum",
			"Liter",
			"Bruttobetrag (€)",
			"Eingangsrechnung",
			"Bemerkung",
			*[_label(d) for d in fields],
		],
		rows,
	)
	fields = _extras(definitions, "Kosten")
	sheet(
		"Heizungsnebenkosten",
		[
			"Bezeichnung",
			"Belegdatum",
			"Bruttobetrag (€)",
			"Eingangsrechnung",
			"Bemerkung",
			*[_label(d) for d in fields],
		],
		[
			[
				r.bezeichnung,
				_date(r.datum),
				r.betrag if r.betrag_bestaetigt else None,
				r.eingangsrechnung,
				r.bemerkung,
				*[_value(d, json_object(r.zusatzwerte_json)) for d in fields],
			]
			for r in doc.kosten
		],
	)
	values = json_object(doc.zusatzwerte_json)
	sheet(
		"Zusatzangaben",
		["Abschnitt", "Angabe", "Wert", "Einheit", "Ausfüllhinweis"],
		[
			[d.get("abschnitt"), d["bezeichnung"], _value(d, values), d.get("einheit"), d.get("hinweis")]
			for d in _extras(definitions, "Meldung")
		],
		{2: 65, 3: 28, 5: 65},
	)
	if int(doc.docstatus) == 0:
		sheet("Offene Angaben", ["Vor Freigabe zu erledigen"], [[s] for s in doc.issues()], {1: 110})
	stream = BytesIO()
	wb.save(stream)
	return stream.getvalue()

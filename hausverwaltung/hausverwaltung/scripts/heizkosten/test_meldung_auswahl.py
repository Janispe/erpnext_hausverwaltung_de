import unittest
from io import BytesIO
from types import SimpleNamespace as Row

from openpyxl import load_workbook

from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_export import (
	DEFAULT_NUTZER_FIELDS,
	build_xlsx,
	nutzer_columns,
)
from hausverwaltung.hausverwaltung.scripts.heizkosten.test_meldung_v2 import sample


def sample_with_users():
	doc = sample()
	doc.nutzer = [
		Row(
			wohnung_id="18",
			wohnung="W1",
			nutzernummer="016",
			typ=kind,
			mietvertrag=f"MV-{i}" if kind == "Mietvertrag" else None,
			mietername=name,
			von="2025-01-01",
			bis="2025-12-31",
			heizflaeche=40,
			wohnflaeche=42,
			flaeche_bestaetigt=0,
			vorauszahlung_meldung=999,
			vorauszahlung_bestaetigt=0,
			vorauszahlung_ist=paid,
			vorauszahlung_soll=600,
			pruefhinweise="Rechnung prüfen",
			pruefnotiz=None,
			zusatzwerte_json='{"personen": 0}',
		)
		for i, (kind, name, paid) in enumerate(
			[
				("Mietvertrag", "Mieter A", 500),
				("Leerstand", "Leerstand", 0),
				("Mietvertrag", "=1+1", 0),
			]
		)
	]
	return doc


class TestMeldungAuswahl(unittest.TestCase):
	def test_default_list_uses_actual_payments_and_living_area_per_contract(self):
		wb = load_workbook(
			BytesIO(build_xlsx(sample_with_users(), nutzer_fields=list(DEFAULT_NUTZER_FIELDS)))
		)
		self.assertEqual(wb.sheetnames, ["Mieterliste"])
		ws = wb.active
		self.assertEqual(list(ws.values)[4:], [(18, "Mieter A", 500, 42), (18, "=1+1", 0, 42)])
		self.assertTrue(all(ws.cell(row, 1).data_type == "n" for row in range(5, 7)))
		self.assertEqual(ws["A5"].number_format, "0")
		self.assertEqual(ws["B6"].data_type, "s")
		self.assertEqual(ws["D5"].number_format, "#,##0.000")
		self.assertEqual(ws["C5"].number_format, "#,##0.00")

	def test_selected_order_extras_dates_and_vacancies(self):
		doc = sample_with_users()
		fields = [
			"von",
			"zusatz:personen",
			"mietername",
			"wohnung_id",
			"heizflaeche",
			"vorauszahlung_meldung",
		]
		ws = load_workbook(BytesIO(build_xlsx(doc, nutzer_fields=fields, include_vacancies=True))).active
		self.assertEqual(ws.max_row, 7)
		self.assertEqual(ws["A5"].value.date().isoformat(), "2025-01-01")
		self.assertEqual(ws["B5"].value, 0)
		self.assertEqual(ws["C6"].value, "Leerstand")
		self.assertEqual(ws["D5"].value, 18)
		self.assertIsNone(ws["E5"].value)  # unconfirmed heating area
		self.assertIsNone(ws["F5"].value)  # unconfirmed reporting amount
		self.assertEqual(ws.auto_filter.ref, "A4:F7")

	def test_invalid_or_duplicate_columns_rejected(self):
		for fields in (
			[],
			"mietername",
			{},
			["customer"],
			["__dict__"],
			[None],
			[["wohnung"]],
			["wohnung", "wohnung"],
		):
			with self.subTest(fields=fields), self.assertRaises(ValueError):
				build_xlsx(sample_with_users(), nutzer_fields=fields)

	def test_full_export_retains_all_sheets_and_vacancies(self):
		doc = sample_with_users()
		wb = load_workbook(BytesIO(build_xlsx(doc)))
		self.assertIn("Heizoel", wb.sheetnames)
		self.assertNotIn("Mieterliste", wb.sheetnames)
		self.assertEqual(wb["Nutzer"].max_row, 7)
		self.assertEqual(wb["Nutzer"].max_column, len(nutzer_columns(doc)))

	def test_single_column_and_empty_tenant_list(self):
		doc = sample_with_users()
		doc.nutzer = [doc.nutzer[1]]
		ws = load_workbook(BytesIO(build_xlsx(doc, nutzer_fields=["wohnung_id"]))).active
		self.assertEqual(ws.max_row, 4)
		self.assertEqual(ws.auto_filter.ref, "A4:A4")

	def test_apartment_numbers_are_sorted_numerically_and_tenant_changes_stay_chronological(self):
		doc = sample_with_users()
		doc.nutzer[0].wohnung_id = "10"
		doc.nutzer[0].von = "2025-07-01"
		doc.nutzer[1].wohnung_id = "2"
		doc.nutzer[2].wohnung_id = "10"
		doc.nutzer[2].von = "2025-01-01"
		doc.nutzer.append(Row(**{**vars(doc.nutzer[0]), "wohnung_id": "1", "mietername": "Mieter 1"}))
		ws = load_workbook(
			BytesIO(
				build_xlsx(doc, nutzer_fields=["wohnung_id", "mietername", "von"], include_vacancies=True)
			)
		).active
		self.assertEqual([ws.cell(row, 1).value for row in range(5, 9)], [1, 2, 10, 10])
		self.assertTrue(all(ws.cell(row, 1).number_format == "0" for row in range(5, 9)))
		self.assertEqual(
			[ws.cell(row, 3).value.date().isoformat() for row in range(7, 9)], ["2025-01-01", "2025-07-01"]
		)

import base64
import unittest
from io import BytesIO
from types import SimpleNamespace as Row
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from hausverwaltung.hausverwaltung.patches.post_model_sync.seed_heizkostenmeldung_vorlage_v2 import (
	build_fields,
)
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_export import build_xlsx
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_schema import dumps, normalize_definitions
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_summen import summary
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_unterschrift import signature_bytes


def sample():
	definitions = normalize_definitions(build_fields())
	return Row(
		name="TEST-V2",
		docstatus=0,
		immobilie="Test",
		von="2025-01-01",
		bis="2025-12-31",
		liegenschaftsnummer="00123",
		objektadresse="Teststraße 1",
		waermedienst=None,
		datenstand=None,
		notizen=None,
		vorlage_snapshot=dumps({"name": "ares-v2"}),
		zusatzwerte_json=dumps(
			{
				"verwaltung_strasse": "Teststraße 1",
				"verwaltung_plz": "01234",
				"konto_iban": "00123",
				"weitere_hausbewohner": "Zusätzliche Person",
			}
		),
		nutzer=[],
		anfangsbestand=100,
		anfangswert=150,
		endbestand=20,
		endwert=30,
		endwert_durch_messdienst=0,
		abzuege=5,
		soforthilfe=10,
		preisbremse=15,
		bestaende_bestaetigt=1,
		brennstoff_vollstaendig=1,
		kosten_vollstaendig=1,
		lieferungen=[
			Row(
				datum="2025-02-01",
				menge=100,
				betrag=120,
				betrag_bestaetigt=1,
				eingangsrechnung=None,
				bemerkung=None,
				zusatzwerte_json="{}",
			)
		],
		kosten=[
			Row(
				bezeichnung="Wartung",
				datum="2025-03-01",
				betrag=40,
				betrag_bestaetigt=1,
				eingangsrechnung=None,
				bemerkung=None,
				zusatzwerte_json="{}",
			)
		],
		abgaben=[
			Row(
				bezeichnung="Umsatzsteuer",
				bruttobetrag=19,
				mwst_satz=19,
				behandlung="Im Brennstoffbetrag enthalten",
				angaben_bestaetigt=1,
				bemerkung=None,
			),
			Row(
				bezeichnung="Weitere Abgabe",
				bruttobetrag=8,
				mwst_satz=0,
				behandlung="Zusätzlich berechnen",
				angaben_bestaetigt=1,
				bemerkung=None,
			),
		],
		definitions=lambda: definitions,
		issues=lambda: [],
	)


def totals(doc):
	return {r["bezeichnung"]: r["wert"] for r in summary(doc)}


class TestMeldungV2(unittest.TestCase):
	def test_export_keeps_apartment_number_across_tenant_change_and_vacancy(self):
		doc = sample()
		doc.nutzer = [
			Row(
				wohnung_id="123",
				wohnung="W1",
				nutzernummer=None,
				typ="Mietvertrag" if contract else "Leerstand",
				mietvertrag=contract,
				mietername=contract or "Leerstand",
				von="2025-01-01",
				bis="2025-12-31",
				heizflaeche=42,
				wohnflaeche=42,
				flaeche_bestaetigt=1,
				vorauszahlung_meldung=0,
				vorauszahlung_bestaetigt=1,
				vorauszahlung_ist=0,
				vorauszahlung_soll=0,
				pruefhinweise=None,
				pruefnotiz=None,
				zusatzwerte_json="{}",
			)
			for contract in ("MV-A", None, "MV-B")
		]
		ws = load_workbook(BytesIO(build_xlsx(doc)))["Nutzer"]
		self.assertEqual(ws["A4"].value, "Wohnungsnummer (ERP)")
		self.assertEqual([ws.cell(r, 1).value for r in range(5, 8)], ["123"] * 3)
		self.assertTrue(all(ws.cell(r, 1).data_type == "s" for r in range(5, 8)))
		self.assertEqual([ws.cell(r, 2).value for r in range(5, 8)], [None] * 3)
		self.assertEqual([ws.cell(r, ws.max_column).value for r in range(5, 8)], ["MV-A", None, "MV-B"])
		self.assertEqual(ws["H5"].number_format, "#,##0.000")
		self.assertEqual(ws["L5"].number_format, "#,##0.000")

	def test_deductions_and_taxes_are_counted_exactly_once(self):
		doc = sample()
		total = totals(doc)
		self.assertEqual(total["Heizölverbrauch"], 180)
		self.assertEqual(total["Summe Abzüge"], 30)
		self.assertEqual(total["Brennstoffkosten nach Abzügen"], 210)
		self.assertEqual(total["Abzurechnende Gesamtkosten"], 258)
		doc.endwert_durch_messdienst = 1
		self.assertIsNone(totals(doc)["Abzurechnende Gesamtkosten"])

	def test_missing_confirmations_do_not_create_false_totals(self):
		for part in ("lieferungen", "kosten", "abgaben"):
			doc = sample()
			row = getattr(doc, part)[0]
			setattr(row, "angaben_bestaetigt" if part == "abgaben" else "betrag_bestaetigt", 0)
			self.assertIsNone(totals(doc)["Abzurechnende Gesamtkosten"])
		doc = sample()
		doc.lieferungen = doc.kosten = doc.abgaben = []
		doc.anfangswert = doc.endwert = doc.abzuege = doc.soforthilfe = doc.preisbremse = 0
		self.assertEqual(totals(doc)["Abzurechnende Gesamtkosten"], 0)

	def test_turnover_does_not_duplicate_area_but_adds_payments(self):
		doc = sample()
		doc.nutzer = [
			Row(
				wohnung="W1",
				typ="Mietvertrag",
				heizflaeche=50,
				flaeche_bestaetigt=1,
				vorauszahlung_meldung=v,
				vorauszahlung_bestaetigt=1,
				zusatzwerte_json=dumps(
					{"nk_vorauszahlung": n, "nebenkostenflaeche": 50, "warmwasserflaeche": 45}
				),
			)
			for v, n in ((100, 20), (200, 30))
		]
		total = totals(doc)
		self.assertEqual(total["Summe Heizflächen"], 50)
		self.assertEqual(total["Summe Warmwasserflächen"], 45)
		self.assertEqual(total["Summe Heizkostenvorauszahlungen"], 300)
		self.assertEqual(total["Summe Hausnebenkostenvorauszahlungen"], 50)
		doc.nutzer[1].heizflaeche = 60
		self.assertIsNone(totals(doc)["Summe Heizflächen"])

	def test_v2_exports_tax_rates_contact_details_and_confirmation(self):
		doc = sample()
		doc.bestaetigung_ort = "Berlin"
		doc.bestaetigung_datum = "2026-09-09"
		doc.unterzeichner = "Testperson"
		png = BytesIO()
		Image.new("RGB", (40, 20), "black").save(png, format="PNG")
		doc.unterschrift = "data:image/png;base64," + base64.b64encode(png.getvalue()).decode()
		data = build_xlsx(doc)
		wb = load_workbook(BytesIO(data))
		self.assertEqual(wb["Brennstoffabgaben"]["C6"].value, 0)
		self.assertEqual(wb["Bestaetigung"]["C5"].value, "Berlin")
		self.assertEqual(wb["Bestaetigung"]["C6"].value.date().isoformat(), "2026-09-09")
		self.assertEqual(len(wb["Bestaetigung"]._images), 1)
		with ZipFile(BytesIO(data)) as archive:
			self.assertTrue(any(n.startswith("xl/media/") for n in archive.namelist()))
		extras = {r[1].value: r[2].value for r in wb["Zusatzangaben"].iter_rows(min_row=5)}
		self.assertEqual(extras["Verwaltung: Postleitzahl"], "01234")
		self.assertEqual(extras["Kontonummer / IBAN"], "00123")
		self.assertEqual(extras["Weitere Hausbewohner"], "Zusätzliche Person")
		self.assertFalse(any(c.data_type == "f" for s in wb for r in s for c in r))

	def test_signature_rejects_urls_and_invalid_images(self):
		for value in (
			"https://example.com/a.png",
			"data:image/png;base64,aGVsbG8=",
			"data:image/svg+xml;base64,PHN2Zz4=",
			42,
		):
			with self.subTest(value=value), self.assertRaises(ValueError):
				signature_bytes(value)

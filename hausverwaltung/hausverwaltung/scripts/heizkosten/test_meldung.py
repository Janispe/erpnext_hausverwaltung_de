import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_daten import segments
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_export import build_xlsx
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_schema import (
	carry_values,
	dumps,
	initial_values,
	normalize_definitions,
	normalize_values,
)


def definition(**kwargs):
	return normalize_definitions(
		[
			dict(
				schluessel="wert",
				bezeichnung="Wert",
				bereich="Meldung",
				feldtyp="Currency",
				pflichtfeld=1,
				**kwargs,
			)
		]
	)[0]


class TestMeldungSchema(unittest.TestCase):
	def test_default_zero_survives_definition_normalization(self):
		d = definition(standardwert=0)
		self.assertEqual(initial_values([d], "Meldung"), {"wert": 0.0})

	def test_new_version_does_not_reinterpret_old_values(self):
		d = dict(definition(folgejahr_uebernehmen=1), feldtyp="Select", optionen="Ja\nNein")
		changed = dict(d, optionen="Offen\nErledigt")
		self.assertEqual(carry_values([d], [changed], "Meldung", {"wert": "Nein"}), {})
		self.assertEqual(carry_values([d], [dict(d, feldtyp="Float")], "Meldung", {"wert": "Nein"}), {})

	def test_required_zero_is_present_but_blank_is_not(self):
		d = definition()
		self.assertEqual(normalize_values([d], "Meldung", {"wert": 0}, required=True), ({"wert": 0.0}, []))
		self.assertEqual(normalize_values([d], "Meldung", {"wert": ""}, required=True), ({}, ["Wert"]))

	def test_rejects_unknown_fields_and_nonfinite_values(self):
		for value in ({"unknown": 1}, {"wert": "NaN"}, {"wert": "Infinity"}, {"wert": []}, {"wert": True}):
			with self.subTest(value=value), self.assertRaises(ValueError):
				normalize_values([definition()], "Meldung", value)

	def test_rejects_duplicate_keys_and_labels(self):
		d = definition()
		for other in (d, dict(d, schluessel="anderer"), dict(d, bezeichnung="Andere")):
			with self.assertRaises(ValueError):
				normalize_definitions([d, other])

	def test_type_validation_and_versioned_options(self):
		for typ, good, bad in (
			("Date", "2024-02-29", "2025-02-29"),
			("Int", 0, 1.5),
			("Select", "Nein", "Vielleicht"),
		):
			d = dict(definition(), feldtyp=typ, optionen="Ja\nNein")
			self.assertFalse(normalize_values([d], "Meldung", {"wert": good}, required=True)[1])
			with self.assertRaises(ValueError):
				normalize_values([d], "Meldung", {"wert": bad})

	def test_next_year_copies_only_explicit_stable_fields(self):
		d = definition(folgejahr_uebernehmen=1)
		volatile = dict(d, schluessel="kosten", folgejahr_uebernehmen=0)
		self.assertEqual(initial_values([d, volatile], "Meldung", {"wert": 0, "kosten": 500}), {"wert": 0})


class TestMeldungSegments(unittest.TestCase):
	def test_change_and_vacancy_cover_each_day_once(self):
		contracts = [
			dict(name="A", wohnung="W", kunde="C-A", von="2025-01-01", bis="2025-03-15"),
			dict(name="B", wohnung="W", kunde="C-B", von="2025-04-01", bis=None),
		]
		rows = segments(["W"], contracts, "2025-01-01", "2025-12-31")
		self.assertEqual([r["typ"] for r in rows], ["Mietvertrag", "Leerstand", "Mietvertrag"])
		self.assertEqual(rows[1]["von"], date(2025, 3, 16))
		self.assertEqual(rows[1]["bis"], date(2025, 3, 31))
		self.assertEqual(sum((r["bis"] - r["von"]).days + 1 for r in rows), 365)
		self.assertEqual(len({r["zeilen_id"] for r in rows}), 3)
		self.assertEqual([r["customer"] for r in rows], ["C-A", None, "C-B"])

	def test_overlap_or_missing_customer_stops_instead_of_guessing(self):
		a = dict(name="A", wohnung="W", kunde="C-A", von="2025-01-01", bis="2025-03-15")
		b = dict(name="B", wohnung="W", kunde="C-B", von="2025-03-15", bis=None)
		with self.assertRaises(ValueError):
			segments(["W"], [a, b], "2025-01-01", "2025-12-31")
		with self.assertRaises(ValueError):
			segments(["W"], [dict(a, kunde=None)], "2025-01-01", "2025-12-31")

	def test_empty_apartment_and_leap_year(self):
		rows = segments(["W"], [], "2024-01-01", "2024-12-31")
		self.assertEqual(rows[0]["typ"], "Leerstand")
		self.assertEqual((rows[0]["bis"] - rows[0]["von"]).days + 1, 366)


class TestMeldungExport(unittest.TestCase):
	def test_snapshot_extra_fields_dates_zero_and_excel_formula_injection(self):
		field = dict(definition(), feldtyp="Data", bezeichnung="=1+1")
		doc = SimpleNamespace(
			name="HKM-TEST",
			docstatus=0,
			immobilie='=HYPERLINK("x")',
			von="2025-01-01",
			bis="2025-12-31",
			liegenschaftsnummer="00123",
			objektadresse="Test",
			waermedienst=None,
			datenstand=None,
			notizen=None,
			vorlage_snapshot=dumps({"name": "ares-v7", "felder": [field]}),
			zusatzwerte_json=dumps({"wert": '=WEBSERVICE("x")'}),
			nutzer=[],
			lieferungen=[],
			kosten=[],
			bestaende_bestaetigt=1,
			brennstoff_vollstaendig=1,
			anfangsbestand=0,
			anfangswert=0,
			endbestand=0,
			endwert=0,
			endwert_durch_messdienst=1,
			abzuege=0,
			definitions=lambda: [field],
			issues=lambda: ["Daten fehlen"],
		)
		wb = load_workbook(BytesIO(build_xlsx(doc)))
		self.assertEqual(wb["Zusatzangaben"]["C5"].value, '=WEBSERVICE("x")')
		self.assertEqual(wb["Zusatzangaben"]["C5"].data_type, "s")
		self.assertEqual(wb["Meldung"]["D2"].value, "00123")
		self.assertEqual(wb["Heizoel"]["C5"].value, 0)
		self.assertEqual(wb["Heizoel"]["B5"].value.date(), date(2025, 1, 1))
		self.assertIn("Offene Angaben", wb.sheetnames)
		for sheet in wb:
			self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))

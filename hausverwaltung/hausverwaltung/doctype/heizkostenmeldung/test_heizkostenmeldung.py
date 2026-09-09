"""Frappe-Lebenszyklus-Tests; nur auf einer explizit erlaubten Test-/Entwicklersite."""

import hashlib
import unittest
from io import BytesIO
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from hausverwaltung.hausverwaltung.doctype.heizkostenmeldung import heizkostenmeldung as module
from hausverwaltung.hausverwaltung.doctype.heizkostenmeldung_vorlage.heizkostenmeldung_vorlage import (
	neue_version,
)
from hausverwaltung.hausverwaltung.scripts.heizkosten.meldung_schema import dumps, json_object


class TestHeizkostenmeldung(unittest.TestCase):
	def setUp(self):
		if not getattr(frappe.local, "site", None):
			self.skipTest("Frappe-Testsite erforderlich")
		self.user = frappe.session.user
		frappe.set_user("Administrator")
		self.point = "hk_meldung_test"
		frappe.db.savepoint(self.point)
		self.files = []
		self.property = frappe.db.get_value("Immobilie", {}, "name")
		if not self.property:
			self.skipTest("Testsite benötigt eine Immobilie")
		self.template = frappe.get_doc(
			{
				"doctype": "Heizkostenmeldung Vorlage",
				"vorlagenkennung": "test_" + frappe.generate_hash(length=10),
				"version": 1,
				"bezeichnung": "Testvorlage",
				"felder": [
					{
						"schluessel": "zahl",
						"bezeichnung": "Zahl",
						"bereich": "Meldung",
						"feldtyp": "Currency",
						"pflichtfeld": 1,
					},
					{
						"schluessel": "kontakt",
						"bezeichnung": "Kontakt",
						"bereich": "Meldung",
						"feldtyp": "Data",
						"folgejahr_uebernehmen": 1,
					},
				],
			}
		).insert()
		self.template.submit()
		self.doc = frappe.get_doc(
			{
				"doctype": "Heizkostenmeldung",
				"immobilie": self.property,
				"von": "2025-01-01",
				"bis": "2025-12-31",
				"vorlage": self.template.name,
			}
		).insert()

	def tearDown(self):
		for name in self.files:
			if frappe.db.exists("File", name):
				frappe.delete_doc("File", name, ignore_permissions=True)
		if hasattr(self, "point"):
			frappe.db.rollback(save_point=self.point)
		frappe.set_user(self.user)

	def test_draft_template_cannot_be_used(self):
		name = neue_version(self.template.name)
		draft = frappe.get_doc("Heizkostenmeldung Vorlage", name)
		self.assertEqual(draft.docstatus, 0)
		self.assertEqual(draft.version, 2)
		with self.assertRaises(frappe.ValidationError):
			frappe.get_doc(
				{
					"doctype": "Heizkostenmeldung",
					"immobilie": self.property,
					"von": "2026-01-01",
					"bis": "2026-12-31",
					"vorlage": name,
				}
			).insert()

	def test_new_version_does_not_change_old_snapshot(self):
		original = self.doc.vorlage_snapshot
		draft = frappe.get_doc("Heizkostenmeldung Vorlage", neue_version(self.template.name))
		draft.felder[0].bezeichnung = "Neue Bezeichnung"
		draft.save().submit()
		self.doc.reload()
		self.assertEqual(self.doc.vorlage_snapshot, original)
		self.template.felder[0].bezeichnung = "Manipuliert"
		with self.assertRaises(frappe.ValidationError):
			self.template.save()

	def test_used_template_cannot_be_cancelled(self):
		with self.assertRaises(frappe.ValidationError):
			self.template.cancel()

	def test_client_cannot_forge_snapshot_or_source_refresh_flag(self):
		self.doc.vorlage_snapshot = dumps({"felder": []})
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		self.doc.reload()
		self.doc.flags.hk_refresh = True
		self.doc.objektadresse = "Gefälscht"
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()

	def test_unknown_fields_rejected_and_required_zero_accepted(self):
		self.doc.zusatzwerte_json = dumps({"zahl": 0, "unbekannt": 1})
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		self.doc.reload()
		self.doc.zusatzwerte_json = dumps({"zahl": 0})
		self.doc.save()
		self.assertNotIn("Meldung: Zahl fehlt.", self.doc.issues())

	def test_browser_numeric_roundtrip_preserves_source_protection(self):
		stored = dict(zeilen_id="row", wohnflaeche=101.0, vorauszahlung_ist=1680.0)
		browser = dict(stored, wohnflaeche=101, vorauszahlung_ist=1680)
		self.assertFalse(module._sources_changed(browser, stored))
		self.assertTrue(module._sources_changed(dict(browser, vorauszahlung_ist=1681), stored))
		self.assertTrue(module._sources_changed(dict(browser, customer="another"), stored))
		self.assertTrue(module._sources_changed(dict(browser, wohnung_id="999"), stored))

	def test_internal_apartment_number_is_sufficient_without_service_number(self):
		row = self.doc.append("nutzer", {"wohnung": "W1", "wohnung_id": "123", "typ": "Leerstand"})
		self.assertFalse(any("Wohnungsnummer fehlt" in issue for issue in self.doc.issues()))
		row.wohnung_id = None
		self.assertTrue(any("Wohnungsnummer fehlt" in issue for issue in self.doc.issues()))
		row.nutzernummer = "016"
		self.assertFalse(any("Wohnungsnummer fehlt" in issue for issue in self.doc.issues()))

	def test_required_inputs_block_submit(self):
		with self.assertRaises(frappe.ValidationError):
			self.doc.submit()

	def test_following_year_accepts_new_version_without_copying_year_costs(self):
		self.doc.zusatzwerte_json = dumps({"zahl": 999, "kontakt": "Verwaltung"})
		self.doc.endbestand = 20
		self.doc.endwert = 50
		self.doc.endwert_durch_messdienst = 0
		self.doc.soforthilfe = 10
		self.doc.preisbremse = 20
		self.doc.bestaetigung_datum = "2025-12-31"
		self.doc.unterzeichner = "Testperson"
		self.doc.append(
			"abgaben",
			{
				"bezeichnung": "Umsatzsteuer",
				"bruttobetrag": 19,
				"mwst_satz": 19,
				"angaben_bestaetigt": 1,
				"behandlung": "Im Brennstoffbetrag enthalten",
			},
		)
		self.doc.save()
		next_template = frappe.get_doc("Heizkostenmeldung Vorlage", neue_version(self.template.name))
		next_template.append(
			"felder",
			{
				"schluessel": "neu",
				"bezeichnung": "Neu",
				"bereich": "Meldung",
				"feldtyp": "Data",
				"pflichtfeld": 1,
			},
		)
		next_template.save().submit()
		with patch.object(module, "daten_laden"):
			name = module.folgejahr(self.doc.name, next_template.name)
		next_doc = frappe.get_doc("Heizkostenmeldung", name)
		self.assertEqual(str(next_doc.von), "2026-01-01")
		self.assertEqual(str(next_doc.bis), "2026-12-31")
		self.assertEqual(next_doc.vorlage, next_template.name)
		self.assertEqual(next_doc.anfangsbestand, 20)
		self.assertEqual(next_doc.anfangswert, 50)
		self.assertFalse(next_doc.bestaende_bestaetigt)
		self.assertFalse(next_doc.lieferungen)
		self.assertFalse(next_doc.abgaben)
		self.assertFalse(next_doc.soforthilfe)
		self.assertFalse(next_doc.preisbremse)
		self.assertFalse(next_doc.bestaetigung_datum)
		self.assertFalse(next_doc.unterzeichner)
		self.assertEqual(json_object(next_doc.zusatzwerte_json), {"kontakt": "Verwaltung"})

	def test_abgaben_and_deductions_validation(self):
		self.doc.soforthilfe = -1
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		self.doc.reload()
		self.doc.append(
			"abgaben", {"bezeichnung": "Abgabe", "mwst_satz": 101, "behandlung": "Zusätzlich berechnen"}
		)
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		self.doc.reload()
		self.doc.append(
			"abgaben", {"bezeichnung": "Abgabe", "mwst_satz": 0, "behandlung": "Zusätzlich berechnen"}
		)
		self.doc.save()
		self.assertTrue(any("MwSt.-Satz prüfen" in s for s in self.doc.issues()))
		self.doc.abgaben[0].angaben_bestaetigt = 1
		self.doc.save()
		self.assertFalse(any("MwSt.-Satz prüfen" in s for s in self.doc.issues()))

	def test_v2_seed_preserves_published_v1_and_is_idempotent(self):
		from hausverwaltung.hausverwaltung.patches.post_model_sync.seed_heizkostenmeldung_vorlage_v2 import (
			execute,
		)

		before = frappe.get_doc("Heizkostenmeldung Vorlage", "ares-v1").as_json()
		execute()
		first = frappe.get_doc("Heizkostenmeldung Vorlage", "ares-v2").as_json()
		execute()
		self.assertEqual(frappe.get_doc("Heizkostenmeldung Vorlage", "ares-v1").as_json(), before)
		self.assertEqual(frappe.get_doc("Heizkostenmeldung Vorlage", "ares-v2").as_json(), first)

	def test_changed_occupancy_blocks_release(self):
		with (
			patch.object(module.Heizkostenmeldung, "issues", return_value=[]),
			patch.object(module, "load_segments", return_value=[{"zeilen_id": "new-contract"}]),
		):
			with self.assertRaisesRegex(frappe.ValidationError, "Mietvertragsbelegung"):
				self.doc.submit()

	def test_unauthorized_read_export_refresh_and_version_copy(self):
		frappe.set_user("Guest")
		for fn, name in (
			(module.pruefen, self.doc.name),
			(module.export_xlsx, self.doc.name),
			(module.daten_laden, self.doc.name),
			(module.summen, self.doc.name),
			(neue_version, self.template.name),
		):
			with self.subTest(fn=fn.__name__), self.assertRaises(frappe.PermissionError):
				fn(name)

	def test_archive_is_private_immutable_and_download_is_same_bytes(self):
		self.doc.zusatzwerte_json = dumps({"zahl": 0})
		self.doc.save()
		# Isolate File/Document lifecycle from source data; completeness is tested separately.
		with (
			patch.object(module.Heizkostenmeldung, "before_submit"),
			patch.object(module.Heizkostenmeldung, "issues", return_value=[]),
		):
			self.doc.submit()
		file = frappe.get_doc("File", {"file_url": self.doc.export_datei, "attached_to_name": self.doc.name})
		self.files.append(file.name)
		self.assertTrue(file.is_private)
		content = file.get_content()
		self.assertEqual(hashlib.sha256(content).hexdigest(), self.doc.export_sha256)
		self.assertNotIn("Offene Angaben", load_workbook(BytesIO(content)).sheetnames)
		module.export_xlsx(self.doc.name)
		self.assertEqual(frappe.local.response.filecontent, content)
		self.doc.anfangsbestand = 100
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		self.doc.reload()
		self.doc.versandt_von = "Administrator"
		with self.assertRaises(frappe.ValidationError):
			self.doc.save()
		module.versand_vermerken(self.doc.name)
		self.doc.reload()
		self.assertTrue(self.doc.versandt_am)
		self.assertEqual(self.doc.versandt_von, "Administrator")
		module.export_xlsx(self.doc.name)
		self.assertEqual(frappe.local.response.filecontent, content)
